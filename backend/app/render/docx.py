"""Word (.docx) renderer for the document outline.

Matches the team's Creative Brief template: Arial 11, bold title, "BRAND:" line, numbered blue headings,
bullets of **Label:** value, a grey rule between sections, US Letter portrait with 1" margins. The file opens
in Google Docs (Drive upload or File → Open), Word and Pages with the same structure.
"""
from __future__ import annotations

import io

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from ..house_format import FormatSpec
from ..schemas import Job
from . import document as doc


def _rgb(hex_: str) -> RGBColor:
    h = hex_.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


class _Doc:
    def __init__(self, spec: FormatSpec):
        self.st = spec.style
        self.d = Document()
        sec = self.d.sections[0]
        w, h = Inches(8.5), Inches(11)
        if self.st.landscape:
            sec.orientation, (w, h) = WD_ORIENT.LANDSCAPE, (h, w)
        sec.page_width, sec.page_height = w, h
        for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(sec, m, Inches(1))
        normal = self.d.styles["Normal"]
        normal.font.name = self.st.font_family
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), self.st.font_family)
        normal.font.size = Pt(self.st.body_pt)
        normal.paragraph_format.space_after = Pt(0)
        normal.paragraph_format.line_spacing = 1.15
        for name in ("List Bullet", "List Bullet 2"):
            style = self.d.styles[name]
            style.paragraph_format.space_after = Pt(1)

    def para(self, text: str = "", *, bold=False, size=None, color=None, before=0, after=0, italic=False):
        p = self.d.add_paragraph()
        p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(before), Pt(after)
        if text:
            self.run(p, text, bold=bold, size=size, color=color, italic=italic)
        return p

    @staticmethod
    def run(p, text, *, bold=False, size=None, color=None, italic=False):
        r = p.add_run(text)
        r.bold, r.italic = bold, italic
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = _rgb(color)
        return r

    def rule(self, color: str, width: int = 6, before: int = 10):
        p = self.para(before=before, after=4)
        pbdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        for k, v in (("w:val", "single"), ("w:sz", str(width)), ("w:space", "1"), ("w:color", color.lstrip("#"))):
            bottom.set(qn(k), v)
        pbdr.append(bottom)
        p._p.get_or_add_pPr().append(pbdr)

    def link(self, p, url: str, text: str):
        rid = p.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                               is_external=True)
        h = OxmlElement("w:hyperlink")
        h.set(qn("r:id"), rid)
        r = OxmlElement("w:r")
        rpr = OxmlElement("w:rPr")
        color = OxmlElement("w:color"); color.set(qn("w:val"), "1155CC")
        u = OxmlElement("w:u"); u.set(qn("w:val"), "single")
        rpr.append(color); rpr.append(u); r.append(rpr)
        t = OxmlElement("w:t"); t.text = text; t.set(qn("xml:space"), "preserve")
        r.append(t); h.append(r); p._p.append(h)

    def items(self, items: list[doc.Item]):
        for it in items:
            p = self.d.add_paragraph(style="List Bullet")
            if it.label:
                self.run(p, f"{it.label}: ", bold=True)
            if isinstance(it.value, list):
                for v in it.value:
                    q = self.d.add_paragraph(style="List Bullet 2")
                    self._text(q, v)
            elif it.muted:
                self.run(p, it.value, color="#999999")
            elif it.link:
                self.link(p, it.link, it.value)
            else:
                self.run(p, it.value)

    def _text(self, p, v: str):
        if v.startswith(("http://", "https://")) and " " not in v:
            self.link(p, v, v)
        else:
            self.run(p, v)

    def table(self, t: doc.Table):
        tb = self.d.add_table(rows=1, cols=len(t.header))
        tb.style = "Table Grid"
        tb.alignment = WD_TABLE_ALIGNMENT.CENTER
        for c, h in zip(tb.rows[0].cells, t.header):
            c.text = ""
            self.run(c.paragraphs[0], h, bold=True, size=self.st.body_pt - 1)
        for row in t.rows:
            for c, v in zip(tb.add_row().cells, row):
                c.text = ""
                self.run(c.paragraphs[0], v, size=self.st.body_pt - 1)


def render_script_docx(job: Job, spec: FormatSpec) -> bytes:
    out = _Doc(spec)
    st = spec.style
    for n, part in enumerate(doc.build(job, spec, include_qc=False)):
        if n:
            out.rule(st.accent_hex, width=12, before=20)
        out.para(part.title, bold=True, size=st.h1_pt, color=st.title_hex, before=12 if n else 0, after=6)
        if part.brand:
            out.para(f"BRAND: {part.brand}", bold=True, size=st.h2_pt, after=2)
        if part.meta:
            out.para(part.meta, size=st.body_pt - 1, color="#666666", after=4)
        for sec in part.sections:
            out.para(f"{sec.number}. {sec.title}", bold=True, size=st.h2_pt, color=st.accent_hex, before=14, after=4)
            for bl in sec.blocks:
                if isinstance(bl, doc.Item):
                    out.items([bl])
                elif isinstance(bl, doc.Sub):
                    out.para(bl.text, bold=True, before=8, after=2)
                    out.items(bl.items)
                elif isinstance(bl, doc.Note):
                    out.para(bl.text, size=st.body_pt - 1, color="#666666", before=2, after=2)
                elif isinstance(bl, doc.Table):
                    out.table(bl)
            out.rule(st.rule_hex)
    buf = io.BytesIO()
    out.d.save(buf)
    return buf.getvalue()
