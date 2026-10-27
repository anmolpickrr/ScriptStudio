"""Brief upload: the team's Creative Brief template (as downloaded from Google Docs) fills the form."""
import base64
import io

from docx import Document

from tests.test_api import H, client  # noqa: F401  (fixture)


def template_docx(values: dict) -> bytes:
    d = Document()
    d.add_heading("Creative Brief", 2)
    d.add_paragraph(f"BRAND: {values.get('brand', '')}")
    sections = {
        "1. Basic Information": ["Product Link", "Logo Image", "Color Palette or Brand Guidelines Link (if available)",
                                 "Raw/Clicked Product Link (Images/Videos)"],
        "2. Targeting & Audience": ["Primary Audience (if available)", "Creative Language ", "Creator Direction "],
        "3. Ad Format & Specifications": ["Platform", "Ad Type"],
        "4. Messaging & Copy": ["Headline Options", "Subheadline / Supporting Copy", "Key USPs (Unique Selling Points)",
                                "Offer (if any)", "Call-To-Action (CTA)"],
        "5. Creative Direction (Summary)": ["Highlight Creative Direction"],
        "6. Visual Direction": ["Design Style", "Reference Links (if available)",
                                "Elements to Highlight (e.g., product, benefits, lifestyle shots)"],
    }
    for head, labels in sections.items():
        d.add_heading(head, 3)
        for lab in labels:
            p = d.add_paragraph(style="List Bullet")
            p.add_run(f"{lab}:").bold = True
            if values.get(lab.strip()):
                p.add_run(" " + values[lab.strip()])
    buf = io.BytesIO(); d.save(buf)
    return buf.getvalue()


def upload(client, data: bytes, name="brief.docx"):
    return client.post("/api/brief/import", headers=H,
                       json={"filename": name, "data_base64": base64.b64encode(data).decode()})


def test_filled_template_maps_to_form_fields(client):
    r = upload(client, template_docx({
        "brand": "Brewly", "Product Link": "https://brewly.in/products/morning-brew",
        "Primary Audience (if available)": "(Age, location, interests) Professionals 22-30, metros",
        "Creative Language": "Hinglish", "Platform": "Meta", "Ad Type": "Video",
        "Headline Options": "Cafe coffee in 30 seconds; Your 9AM saviour",
        "Key USPs (Unique Selling Points)": "100% Arabica; 10 cups per bottle", "Offer (if any)": "Flat 10% off",
        "Call-To-Action (CTA)": "Shop now on brewly.in", "Design Style": "Click to select",
        "Elements to Highlight (e.g., product, benefits, lifestyle shots)": "Pour shot; first sip",
    }))
    assert r.status_code == 200, r.text
    f = r.json()["fields"]
    assert f["product_url"] == "https://brewly.in/products/morning-brew" and f["brand_name"] == "Brewly"
    assert f["target_audience"] == "Professionals 22-30, metros" and f["language"] == "Hinglish"
    assert f["platform"] == "meta_feed" and f["ad_type"] == "Video" and f["offer"] == "Flat 10% off"
    assert f["headline_ideas"] == ["Cafe coffee in 30 seconds", "Your 9AM saviour"]
    assert f["key_usps"] == ["100% Arabica", "10 cups per bottle"] and "design_style" not in f
    assert "Highlight Creative Direction" in f["client_brief"]


def test_bad_uploads_are_explained(client):
    assert "Word" in upload(client, b"not a docx").json()["detail"]
    assert ".docx" in upload(client, b"x", name="brief.pdf").json()["detail"]
    assert "empty template" in upload(client, template_docx({})).json()["detail"]
