"""API-level tests: auth, CSRF, validation, SSRF, ownership, quotas and the full job lifecycle."""
import pytest
from fastapi.testclient import TestClient

from app import auth, store as store_mod, worker
from app.auth import User, make_session
from app.config import settings
from app.ingest import product
from app.llm import FakeLLM
from app.main import app
from app.pipeline import Pipeline
from app.schemas import JobStatus
from app.routes import jobs as jobs_routes
from tests.fixtures import PRODUCT_HTML, handlers

H = {"X-Requested-With": "scriptwright"}
BODY = {"product_url": "https://brewly.in/products/morning-brew",
        "client_brief": "Launch film for young professionals who skip coffee on busy mornings.",
        "platform": "instagram_reels", "duration_sec": 15, "mandatories": ["Shop now on brewly.in"]}


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(store_mod, "_store", store_mod.LocalStore(tmp_path))
    monkeypatch.setattr(settings, "auth_mode", "none")
    monkeypatch.setattr(settings, "daily_job_limit", 15)
    monkeypatch.setattr(jobs_routes, "check_url", lambda u: u)
    monkeypatch.setattr(product, "fetch", lambda url, client=None: product.parse_html(url, PRODUCT_HTML))

    def run_now(task):  # execute synchronously with the fake model instead of a background worker
        monkeypatch.setattr(worker, "get_llm", lambda: FakeLLM(handlers()))
        worker.execute(task)
    monkeypatch.setattr(jobs_routes, "dispatch", run_now)
    return TestClient(app)


def test_post_without_csrf_header_is_rejected(client):
    assert client.post("/api/jobs", json=BODY).status_code == 403


def test_validation_errors_are_readable(client):
    r = client.post("/api/jobs", json={**BODY, "client_brief": "short", "duration_sec": 999}, headers=H)
    assert r.status_code == 422
    assert "client_brief" in r.json()["detail"] and "duration_sec" in r.json()["detail"]


def test_private_urls_are_rejected(client, monkeypatch):
    from app.netsafe import check_url
    monkeypatch.setattr(jobs_routes, "check_url", check_url)
    for url in ["http://127.0.0.1:8080/", "http://169.254.169.254/computeMetadata/v1/", "http://localhost/x",
                "http://metadata.google.internal/"]:
        r = client.post("/api/jobs", json={**BODY, "product_url": url}, headers=H)
        assert r.status_code == 422, url


def test_full_lifecycle(client):
    r = client.post("/api/jobs", json=BODY, headers=H)
    assert r.status_code == 200, r.text
    jid = r.json()["id"]
    job = client.get(f"/api/jobs/{jid}").json()
    assert job["status"] == "ready" and job["script"]["title"]
    assert "source_excerpt" not in job["dossier"]

    html = client.get(f"/api/jobs/{jid}/preview").text
    assert "Scene-by-Scene Script" in html
    assert client.get(f"/api/jobs/{jid}/download").headers["content-disposition"].startswith("attachment")

    assert client.post(f"/api/jobs/{jid}/revise", json={"feedback": "Make the hook punchier."}, headers=H).status_code == 200
    job = client.get(f"/api/jobs/{jid}").json()
    assert job["status"] == "ready" and job["revisions"][0]["kind"] == "feedback"

    assert client.post(f"/api/jobs/{jid}/concept", json={"index": 2}, headers=H).status_code == 200
    assert client.get(f"/api/jobs/{jid}").json()["concept_decision"]["chosen_index"] == 2

    assert client.post(f"/api/jobs/{jid}/approve", json={"notes": "Specific hook"}, headers=H).json()["status"] == "approved"
    assert client.get("/api/library").json()[0]["editor_notes"] == "Specific hook"

    assert client.delete(f"/api/jobs/{jid}", headers=H).status_code == 200
    assert client.get(f"/api/jobs/{jid}").status_code == 404


def test_google_auth_mode_requires_session_and_scopes_jobs(client, monkeypatch):
    monkeypatch.setattr(settings, "auth_mode", "google")
    monkeypatch.setattr(settings, "admin_emails", [])
    assert client.get("/api/jobs").status_code == 401

    alice = make_session(User(email="alice@agency.com"))
    bob = make_session(User(email="bob@agency.com"))
    client.cookies.set(auth.COOKIE, alice)
    jid = client.post("/api/jobs", json=BODY, headers=H).json()["id"]
    assert client.get(f"/api/jobs/{jid}").status_code == 200

    client.cookies.set(auth.COOKIE, bob)
    assert client.get(f"/api/jobs/{jid}").status_code == 404      # can't see someone else's script
    assert client.get("/api/jobs").json() == []
    assert client.put("/api/formats/house-v1", json={}, headers=H).status_code in (403, 422)

    client.cookies.set(auth.COOKIE, alice[:-2] + "00")             # tampered cookie
    assert client.get("/api/jobs").status_code == 401


def test_domain_allowlist(monkeypatch):
    monkeypatch.setattr(settings, "auth_mode", "google")
    monkeypatch.setattr(settings, "allowed_email_domains", ["agency.com"])
    monkeypatch.setattr(settings, "allowed_emails", ["freelancer@gmail.com"])
    assert auth.is_allowed("a@agency.com") and auth.is_allowed("freelancer@gmail.com")
    assert not auth.is_allowed("someone@gmail.com")


def test_daily_quota(client, monkeypatch):
    monkeypatch.setattr(settings, "auth_mode", "google")
    monkeypatch.setattr(settings, "daily_job_limit", 1)
    client.cookies.set(auth.COOKIE, make_session(User(email="carol@agency.com")))
    assert client.post("/api/jobs", json=BODY, headers=H).status_code == 200
    r = client.post("/api/jobs", json=BODY, headers=H)
    assert r.status_code == 429 and "limit" in r.json()["detail"]


def test_task_endpoint_requires_oidc(client, monkeypatch):
    monkeypatch.setattr(settings, "tasks_queue", "projects/p/locations/l/queues/q")
    monkeypatch.setattr(settings, "public_base_url", "https://example.run.app")
    r = client.post("/tasks/run", json={"job_id": "x"})
    assert r.status_code == 401
    r = client.post("/tasks/run", json={"job_id": "x"}, headers={"Authorization": "Bearer forged"})
    assert r.status_code == 401


def test_failed_model_call_marks_job_failed_with_readable_error(tmp_path):
    from app.schemas import Job, JobInputs
    s = store_mod.LocalStore(tmp_path)
    h = handlers()
    h["concepts"] = lambda **_: (_ for _ in ()).throw(RuntimeError("Error code: 429 - rate limit"))
    job = Pipeline(FakeLLM(h), s, page_override=product.parse_html(BODY["product_url"], PRODUCT_HTML)).run(
        Job(id="f1", inputs=JobInputs(**BODY)))
    assert job.status == "failed" and "usage limit" in job.error


def test_security_headers_and_spa_fallback(client):
    r = client.get("/api/health")
    assert r.headers["x-content-type-options"] == "nosniff" and "frame-ancestors" in r.headers["content-security-policy"]
    assert client.get("/api/does-not-exist").status_code == 404


# ---------------------------------------------------------------- guest (anonymous) mode

def test_anonymous_mode_gives_each_browser_a_private_workspace(client, monkeypatch):
    monkeypatch.setattr(settings, "auth_mode", "anonymous")
    cfg = client.get("/api/config").json()
    assert cfg["user"]["guest"] and auth.COOKIE in client.cookies
    jid = client.post("/api/jobs", json=BODY, headers=H).json()["id"]
    assert client.get(f"/api/jobs/{jid}").status_code == 200

    other = TestClient(app)
    other.get("/api/config")
    assert other.get(f"/api/jobs/{jid}").status_code == 404
    assert other.get("/api/jobs").json() == []


def test_admin_passcode_unlocks_admin(client, monkeypatch):
    monkeypatch.setattr(settings, "auth_mode", "anonymous")
    monkeypatch.setattr(settings, "admin_passcode", "open-sesame")
    client.get("/api/config")
    assert client.get("/api/admin/status").status_code == 403
    assert client.post("/api/auth/admin", json={"passcode": "wrong"}, headers=H).status_code == 403
    assert client.post("/api/auth/admin", json={"passcode": "open-sesame"}, headers=H).json()["admin"] is True
    assert client.get("/api/admin/status").status_code == 200


def test_global_daily_cap(client, monkeypatch):
    monkeypatch.setattr(settings, "auth_mode", "anonymous")
    monkeypatch.setattr(settings, "global_daily_limit", 1)
    client.get("/api/config")
    assert client.post("/api/jobs", json=BODY, headers=H).status_code == 200
    other = TestClient(app)
    other.get("/api/config")
    r = other.post("/api/jobs", json=BODY, headers=H)
    assert r.status_code == 429 and "capacity" in r.json()["detail"]


def test_docx_export_follows_house_format(client):
    import io
    from docx import Document
    jid = client.post("/api/jobs", json={**BODY, "offer": "Flat 10% off", "design_style": "Creator-shot UGC"}, headers=H).json()["id"]
    r = client.get(f"/api/jobs/{jid}/docx")
    assert r.status_code == 200 and r.headers["content-disposition"].endswith('.docx"')
    d = Document(io.BytesIO(r.content))
    text = [p.text for p in d.paragraphs]
    assert text[0] == "Creative Brief" and text[1].startswith("BRAND:")
    assert "1. Basic Information" in text and "6. Visual Direction" in text and "5. Scene-by-Scene Script" in text
    assert "Offer: Flat 10% off" in text and "Design Style: Creator-shot UGC" in text
    bullets = [p for p in d.paragraphs if p.style.name == "List Bullet"]
    assert bullets and bullets[0].runs[0].bold and bullets[0].runs[0].text == "Product Link: "
    assert d.sections[0].page_width < d.sections[0].page_height  # portrait, like the template
    assert any(t.startswith("Scene 4") for t in text) and any("Shop now on brewly.in" in t for t in text)


def test_preview_is_responsive(client):
    jid = client.post("/api/jobs", json=BODY, headers=H).json()["id"]
    html = client.get(f"/api/jobs/{jid}/preview").text
    assert "viewport" in html and "sw-stack" in html and "Quality Check Summary" not in html


def test_retry_failed_job(client, monkeypatch):
    jid = client.post("/api/jobs", json=BODY, headers=H).json()["id"]
    assert client.post(f"/api/jobs/{jid}/retry", headers=H).status_code == 409  # not failed
    job = store_mod.get_store().get_job(jid)
    job.status, job.error = JobStatus.failed, "boom"
    store_mod.get_store().save_job(job)
    assert client.post(f"/api/jobs/{jid}/retry", headers=H).status_code == 200
    job = client.get(f"/api/jobs/{jid}").json()
    assert job["status"] == "ready" and job["error"] == ""
    assert sum(1 for f in job["flags"] if "withheld" in f["message"]) == 1  # flags not duplicated by the retry


def test_restart_recovery(client, monkeypatch):
    from app.schemas import Job, JobInputs
    s = store_mod.get_store()
    s.save_job(Job(id="new1", inputs=JobInputs(**BODY), status=JobStatus.running))
    done = client.post("/api/jobs", json=BODY, headers=H).json()["id"]
    j = s.get_job(done)
    j.status = JobStatus.running
    s.save_job(j)
    ran = []
    monkeypatch.setattr(worker, "dispatch", lambda task: ran.append(task.job_id))
    assert worker.recover_interrupted_jobs() == 2
    assert ran == ["new1"] and s.get_job("new1").status == "queued"
    assert s.get_job(done).status == "ready" and "interrupted" in s.get_job(done).error


def test_unreachable_database_gives_clear_errors(monkeypatch):
    monkeypatch.setattr(settings, "store_backend", "postgres")
    monkeypatch.setattr(settings, "database_url", "postgresql://user:secretpw@127.0.0.1:1/db?connect_timeout=2")
    monkeypatch.setattr(store_mod, "_store", None)
    import psycopg_pool
    real = psycopg_pool.ConnectionPool
    monkeypatch.setattr(psycopg_pool, "ConnectionPool", lambda *a, **k: real(*a, **{**k, "timeout": 2}))
    c = TestClient(app)
    deep = c.get("/api/health/deep").json()
    assert deep["ok"] is False and deep["db"].startswith("error:") and "secretpw" not in deep["db"]
    r = c.get("/api/options")
    assert r.status_code == 503 and "DATABASE_URL" in r.json()["detail"]
    monkeypatch.setattr(store_mod, "_store", None)


def test_github_pages_site_uses_header_session_across_origins(client, monkeypatch):
    monkeypatch.setattr(settings, "auth_mode", "anonymous")
    pages = {"Origin": "https://anmolpickrr.github.io"}
    pre = client.options("/api/jobs", headers={**pages, "Access-Control-Request-Method": "POST",
                                               "Access-Control-Request-Headers": "content-type,x-requested-with,x-session-token"})
    assert pre.status_code == 200 and pre.headers["access-control-allow-origin"] == pages["Origin"]
    assert client.options("/api/jobs", headers={"Origin": "https://evil.example", "Access-Control-Request-Method": "POST"}
                          ).headers.get("access-control-allow-origin") is None

    browser = TestClient(app)  # no cookies are sent cross-origin: the token travels in a header instead
    r = browser.get("/api/config", headers=pages)
    token = r.headers["x-session-token"]
    assert "x-session-token" in r.headers["access-control-expose-headers"].lower()
    browser.cookies.clear()
    me = {**pages, **H, "X-Session-Token": token}
    jid = browser.post("/api/jobs", json=BODY, headers=me).json()["id"]
    assert browser.get(f"/api/jobs/{jid}", headers=me).status_code == 200
    assert TestClient(app).get(f"/api/jobs/{jid}", headers=pages).status_code == 404  # another visitor can't see it
